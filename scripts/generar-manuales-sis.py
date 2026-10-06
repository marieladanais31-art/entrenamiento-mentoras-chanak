"""Genera manuales familiares y esquemas de coordinación desde los datos de formación.
No usa capturas privadas ni accede al SIS. python scripts/generar-manuales-sis.py DIRECTORIO
"""
from pathlib import Path
import json
import sys
from html import escape
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.graphics import renderSVG
from docx import Document
from docx.shared import Inches, Pt, RGBColor
import fitz

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / 'public/formacion'
OUT.mkdir(parents=True, exist_ok=True)
NAVY = colors.HexColor('#102d50')
TEAL = colors.HexColor('#087f82')
GOLD = colors.HexColor('#dfad45')
LIGHT = colors.HexColor('#edf3f7')
WIDTH = A4[0] - 100
for name, file in [('Chanak','DejaVuSans.ttf'),('ChanakBold','DejaVuSans-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name, '/usr/share/fonts/truetype/dejavu/' + file))
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyC',fontName='Chanak',fontSize=10,leading=14,textColor=NAVY,spaceAfter=8))
styles.add(ParagraphStyle(name='TitleC',fontName='ChanakBold',fontSize=24,leading=30,textColor=NAVY,spaceAfter=16))
styles.add(ParagraphStyle(name='HeadC',fontName='ChanakBold',fontSize=17,leading=23,textColor=TEAL,spaceAfter=16))
styles.add(ParagraphStyle(name='SmallC',fontName='Chanak',fontSize=8,leading=11,textColor=NAVY,spaceAfter=6))
styles.add(ParagraphStyle(name='CellC',fontName='Chanak',fontSize=8.5,leading=12,textColor=NAVY))

def p(text, style='BodyC'):
    return Paragraph(escape(text), styles[style])

def table(rows):
    cols=len(rows[0]); widths=[WIDTH/cols]*cols
    result=Table([[p(c,'CellC') for c in row] for row in rows], colWidths=widths, hAlign='LEFT', repeatRows=1)
    result.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),LIGHT),('VALIGN',(0,0),(-1,-1),'TOP'),('BOX',(0,0),(-1,-1),.5,colors.HexColor('#cbd5df')),('INNERGRID',(0,0),(-1,-1),.3,colors.HexColor('#dce4eb')),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
    return result

def panel(title, headings, values):
    d=Drawing(WIDTH,170)
    d.add(Rect(0,0,WIDTH,170,fillColor=colors.white,strokeColor=colors.HexColor('#ccd8e2'),strokeWidth=.7))
    d.add(Rect(0,143,WIDTH,27,fillColor=NAVY,strokeColor=NAVY))
    d.add(String(12,153,'CHANAK · '+title,fontName='ChanakBold',fontSize=10,fillColor=colors.white))
    d.add(String(12,128,'ESQUEMA FORMATIVO · DATOS FICTICIOS · NO ES UNA CAPTURA DEL SIS',fontName='ChanakBold',fontSize=7,fillColor=TEAL))
    d.add(Rect(12,95,WIDTH-24,23,fillColor=LIGHT,strokeColor=LIGHT))
    d.add(String(20,103,'ID-DEMO-001   |   Año 2026–2027   |   Periodo de ejemplo',fontName='Chanak',fontSize=8,fillColor=NAVY))
    cw=(WIDTH-24)/len(headings)
    for i,(h,v) in enumerate(zip(headings,values)):
        x=12+i*cw
        d.add(String(x+6,77,h,fontName='ChanakBold',fontSize=8,fillColor=TEAL))
        # Wrap long cells using measured widths.
        line=''; y=58
        for word in v.split():
            nextline=(line+' '+word).strip()
            if pdfmetrics.stringWidth(nextline,'Chanak',7.5)>cw-12 and line:
                d.add(String(x+6,y,line,fontName='Chanak',fontSize=7.5,fillColor=NAVY));y-=11;line=word
            else: line=nextline
        if line:d.add(String(x+6,y,line,fontName='Chanak',fontSize=7.5,fillColor=NAVY))
    d.add(String(18,14,'Comprobar identidad → revisar evidencia → confirmar estado',fontName='Chanak',fontSize=8,fillColor=TEAL))
    return d

def family_panel(kind):
    configs={
      'familia':('Portal de Padres · Mis Hijos',['Estudiante','Accesos','Estado'],['ID-DEMO-001','PEI · Evaluaciones · Pagos','Revisar programa y ciclo']),
      'evidencias':('Evaluaciones & Evidencias',['Evaluación','Evidencia','Estado'],['Unidad de ejemplo','Enlace institucional','En revisión']),
      'dual':('Campus Dual · Estudiante',['Mi progreso','Mis cursos','Documentos / Evidencias'],['Ver estado revisado','Curso asignado A','Atender feedback']),
    }
    title,headings,values=configs[kind]
    return panel(title,headings,values)

def footer(canv,doc):
    canv.saveState()
    canv.setFillColor(NAVY);canv.setFont('Chanak',7)
    canv.drawString(50,27,'Chanak International Academy · Revisión 6 octubre 2026 · '+str(doc.page))
    canv.setStrokeColor(GOLD);canv.line(50,42,A4[0]-50,42)
    canv.restoreState()

def cover(title,subtitle,scope):
    from reportlab.platypus import Image
    logo=Image(str(ROOT/'public/logo-chanak.png'),width=75,height=75)
    logo.hAlign='LEFT'
    return [logo,Spacer(1,32),p('CHANAK INTERNATIONAL ACADEMY','SmallC'),p(title,'TitleC'),p(subtitle),Spacer(1,24),p(scope),Spacer(1,20),p('Nuestra Fe es la Victoria · 1 Juan 5:4'),p('Edición 2026–2027 · Revisada el 6 de octubre de 2026','SmallC'),p('Las imágenes son esquemas explicativos con datos ficticios. Los controles disponibles pueden variar; siga siempre el plan individual y los accesos autorizados.','SmallC'),PageBreak()]

def pdf(name,story,title):
    doc=SimpleDocTemplate(str(OUT/(name+'.pdf')),pagesize=A4,rightMargin=50,leftMargin=50,topMargin=48,bottomMargin=57,title=title,author='Chanak International Academy')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)

def export_family(manual,name):
    story=cover(manual['titulo'],manual['subtitulo'],'Para padres y tutores. Orientación práctica para acompañar al estudiante y consultar su progreso; no sustituye las condiciones de matrícula ni el plan académico aprobado.')
    doc=Document()
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
    for sec in doc.sections:
        sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
        sec.header.paragraphs[0].text='CHANAK INTERNATIONAL ACADEMY · 2026–2027'
        sec.footer.paragraphs[0].text='Revisión 6 octubre 2026 · Uso familiar · Pantallas ilustrativas'
    doc.add_picture(str(ROOT/'public/logo-chanak.png'),width=Inches(1))
    doc.add_heading(manual['titulo'],0);doc.add_paragraph(manual['subtitulo']);doc.add_paragraph('Nuestra Fe es la Victoria · 1 Juan 5:4')
    doc.add_paragraph('Guía por programa. Los esquemas son ilustrativos y contienen datos ficticios. Siga el plan individual y sus accesos autorizados.');doc.add_page_break()
    for page in manual['paginas']:
        story.append(p(page['titulo'],'HeadC'));doc.add_heading(page['titulo'],1)
        if page.get('pantalla'):
            drawing=family_panel(page['pantalla']);story.extend([drawing,Spacer(1,15)])
            svg=renderSVG.drawToString(drawing)
            image=OUT/('esquema_'+page['pantalla']+'.png')
            svg_doc=fitz.open(stream=svg.encode(),filetype='svg')
            rendered=fitz.open(stream=svg_doc.convert_to_pdf(),filetype='pdf')
            rendered[0].get_pixmap(matrix=fitz.Matrix(2,2)).save(image)
            doc.add_picture(str(image),width=Inches(6.1))
        if page.get('tabla'):
            story.extend([table(page['tabla']),Spacer(1,12)])
            rows=page['tabla'];t=doc.add_table(rows=1,cols=len(rows[0]));t.style='Light Shading Accent 1'
            for j,value in enumerate(rows[0]):t.rows[0].cells[j].text=value
            for row in rows[1:]:
                cells=t.add_row().cells
                for j,value in enumerate(row):cells[j].text=value
        for text in page.get('parrafos',[]):story.append(p(text));doc.add_paragraph(text)
        for text in page.get('lista',[]):story.append(p('• '+text));doc.add_paragraph(text,style='List Bullet')
        story.append(PageBreak());doc.add_page_break()
    story.extend([p('Ayuda y actualización','HeadC'),p('Contacto del programa: '+manual['contacto']),p('Cuentas, vínculos y pagos: Administración por el canal institucional recibido.'),p('Para pedir ayuda: sistema, ID, ciclo, pantalla, fecha, mensaje y próxima acción. Nunca envíe contraseñas.'),p('Referencias: guías de acceso Off-Campus y Dual, instructivos de matrícula y Manual de Usuario SIS aportados por Mariela. Accesos y recorridos contrastados con el código actual de SIS, Portal y LMS Dual el 6 octubre 2026. No se ha certificado un nuevo circuito integral autenticado por todos los roles.','SmallC')])
    doc.add_heading('Ayuda y actualización',1);doc.add_paragraph('Contacto del programa: '+manual['contacto']);doc.add_paragraph('Cuentas, vínculos y pagos: Administración por el canal institucional recibido.');doc.add_paragraph('Referencias aportadas por Mariela y recorridos contrastados con el código el 6 octubre 2026. No incluye contraseñas reales.')
    pdf(name,story,manual['titulo']);doc.save(OUT/(name+'.docx'))

manuals=json.loads((ROOT/'src/data/manualesFamilias.json').read_text())
export_family(manuals['offcampus'],'MANUAL_PADRES_OFFCAMPUS_2026-2027')
export_family(manuals['dual'],'MANUAL_PADRES_DUAL_DIPLOMA_2026-2027')
coord=json.loads((ROOT/'src/data/coordinacionSIS.json').read_text())
story=cover(coord['titulo'],coord['subtitulo'],coord['intro'])
for screen in coord['pantallas']:
    story.extend([p(screen['titulo'],'HeadC'),panel(screen['titulo'],screen['columnas'],screen['filas'][0]),Spacer(1,14),p(screen['objetivo']),p('Antes de operar: '+screen['verifica'])])
    for i,text in enumerate(screen['pasos'],1):story.append(p(str(i)+'. '+text))
    story.extend([p('Cierre: '+screen['cierre']),p('Control: '+screen['cuidado']),p('Caso: '+screen['caso']),p('Criterio: '+screen['respuesta']),PageBreak()])
story.append(p(coord['dual']['titulo'],'HeadC'))
for text in coord['dual']['parrafos']:story.append(p(text))
story.append(p('Rutina de coordinación','HeadC'))
for text in coord['rutina']:story.append(p('• '+text))
story.append(PageBreak());story.append(p('Fuentes, límites y validación','HeadC'))
for text in coord['fuentes']:story.append(p(text))
story.extend([p('Las ocho vistas se reconstruyen como esquemas formativos desde el código y la referencia aportada. No son capturas de una sesión autenticada de coordinador; no muestran expedientes ni datos personales reales.'),p('Validación supervisada: comprobar ámbito del hub, filtros, nota/evidencia pendiente, decisión fundamentada, guardado y documento correcto en un entorno de práctica autorizado. No realizar altas, publicar notas ni enviar alertas reales para una demostración.')])
pdf('GUIA_COORDINADOR_PANTALLAS_SIS_2026-2027',story,coord['titulo'])
for f in sorted(OUT.glob('*.pdf')):print(f.name,f.stat().st_size)
